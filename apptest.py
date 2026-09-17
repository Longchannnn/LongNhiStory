import streamlit as st
import pandas as pd
import google.genai as genai
import json
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

st.set_page_config(
    page_title="PDF to Excel Converter",
    page_icon="📄",
    layout="centered"
)

st.title("📄 Trích Xuất PDF Sang Excel (Key-Value)")
st.write("Tải file PDF hóa đơn/debit note để trích xuất dữ liệu thành file Excel theo thứ tự tùy chỉnh.")

# Sidebar để cấu hình API Key
with st.sidebar:
    st.header("⚙️ Cấu hình")
    api_key = st.text_input("Nhập Gemini API Key:", type="password")
    st.markdown("[Lấy Gemini API Key miễn phí tại đây](https://aistudio.google.com/app/apikey)")

# Tải file PDF
uploaded_file = st.file_uploader("Kéo thả hoặc chọn file PDF tại đây", type=["pdf"])

PROMPT_TEXT = """
Bạn là một chuyên gia OCR tài liệu logistics và hóa đơn. Hãy trích xuất toàn bộ dữ liệu từ file PDF này thành dạng danh sách Key-Value (JSON Array).

Cấu trúc JSON yêu cầu trả về:
[
  {"field_name": "Tên trường (giữ nguyên tiếng Anh gốc)", "value": "Giá trị"},
  ...
]

BẮT BUỘC SẮP XẾP THỨ TỰ CÁC TRƯỜNG NHƯ SAU:
1. Thông tin Công ty phát hành (Issuer Company Name, Address, Tel, Fax, Email, Website)
2. Thông tin Chứng từ (Document Type, No., Date)
3. Thông tin Khách hàng (To, Address, Tel/Fax, Taxcode)
4. Thông tin Lô hàng (Job No., POL/AOL, POD/AOD, MBL/MAWB No., CDS No., Q'ty, Container No.)
5. Thông tin TỔNG TIỀN & THANH TOÁN (TOTAL Debit (before tax), Balance due to..., Currency, SAY, Metadata) - ĐẶT ĐOẠN NÀY LÊN TRƯỚC CÁC ITEM
6. Chi tiết từng Item (Item 1 - Category, Description, Q'ty, Unit, Curr., Price, Amount, VAT %, Debit, Credit, Item 2...)

Chú ý:
- Giữ nguyên tên gốc các trường tiếng Anh.
- Chỉ trả về duy nhất chuỗi JSON thuần (JSON Array), không dùng markdown hay ```json.
"""

if uploaded_file:
    if not api_key:
        st.warning("⚠️ Vui lòng nhập Gemini API Key ở thanh bên trái (Sidebar) để tiếp tục.")
    else:
        if st.button("⚡ Bắt đầu trích xuất sang Excel", type="primary"):
            with st.spinner("Đang xử lý PDF và trích xuất dữ liệu bằng AI..."):
                try:
                    # Đọc nội dung PDF
                    pdf_bytes = uploaded_file.read()
                    
                    # Gọi Gemini API
                    client = genai.Client(api_key=api_key)
                    response = client.models.generate_content(
                        model="gemini-2.5-flash",
                        contents=[
                            {"mime_type": "application/pdf", "data": pdf_bytes},
                            PROMPT_TEXT
                        ]
                    )
                    
                    # Parse dữ liệu JSON
                    raw_text = response.text.strip().replace("```json", "").replace("```", "")
                    data_list = json.loads(raw_text)
                    
                    # Hiển thị bảng dữ liệu xem trước
                    df_preview = pd.DataFrame(data_list)
                    df_preview.insert(0, 'STT', range(1, len(df_preview) + 1))
                    df_preview.columns = ["STT", "Field Name / Trường thông tin", "Value / Giá trị"]
                    
                    st.success("✅ Trích xuất thành công!")
                    st.dataframe(df_preview, use_container_width=True)
                    
                    # Xuất Excel có định dạng chuẩn bằng OpenPyXL
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "KeyValue_Data"
                    ws.views.sheetView[0].showGridLines = True

                    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                    font_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
                    font_regular = Font(name="Segoe UI", size=10, color="1E293B")
                    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
                    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                    thin_side = Side(border_style="thin", color="E2E8F0")
                    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

                    headers = ["STT", "Field Name / Trường thông tin", "Value / Giá trị"]
                    for col_idx, h in enumerate(headers, start=1):
                        cell = ws.cell(row=1, column=col_idx, value=h)
                        cell.font = font_header
                        cell.fill = fill_header
                        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else "left", vertical="center")
                        cell.border = border_all
                    ws.row_dimensions[1].height = 24

                    for idx, item in enumerate(data_list, start=1):
                        r = idx + 1
                        field_name = item.get("field_name", "")
                        val = item.get("value", "")

                        c1 = ws.cell(row=r, column=1, value=idx)
                        c1.alignment = Alignment(horizontal="center", vertical="center")
                        c1.font = font_bold
                        c1.border = border_all

                        c2 = ws.cell(row=r, column=2, value=field_name)
                        c2.alignment = Alignment(horizontal="left", vertical="center")
                        c2.font = font_bold
                        c2.border = border_all

                        c3 = ws.cell(row=r, column=3, value=val)
                        c3.font = font_regular
                        c3.border = border_all

                        if isinstance(val, (int, float)):
                            if "VAT" in str(field_name):
                                c3.number_format = '0.0%'
                                c3.alignment = Alignment(horizontal="right", vertical="center")
                            elif "Q'ty" in str(field_name):
                                c3.alignment = Alignment(horizontal="center", vertical="center")
                            else:
                                c3.number_format = '#,##0'
                                c3.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            c3.alignment = Alignment(horizontal="left", vertical="center")

                        if idx % 2 == 0:
                            c1.fill = fill_zebra
                            c2.fill = fill_zebra
                            c3.fill = fill_zebra

                    ws.column_dimensions['A'].width = 10
                    ws.column_dimensions['B'].width = 40
                    ws.column_dimensions['C'].width = 80

                    # Lưu ra bộ nhớ đệm để tạo link tải về
                    excel_buffer = io.BytesIO()
                    wb.save(excel_buffer)
                    excel_data = excel_buffer.getvalue()
                    
                    st.download_button(
                        label="📥 Tải file Excel về máy",
                        data=excel_data,
                        file_name=f"{uploaded_file.name.replace('.pdf', '')}_Key_Value.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary"
                    )
                    
                except Exception as e:
                    st.error(f"❌ Đã xảy ra lỗi: {str(e)}")